import os
import sys
import sqlite3
import csv
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

print("Starting KAKA to Power BI dataset compilation...")

# 1. Paths
base_dir = r"c:\Users\vophi\Downloads\Github\CV\JD-BIM"
p6_db_path = r"C:\Users\vophi\Documents\PPMDBSQLite.db"
boq_sql_path = os.path.join(base_dir, "kaka", "Mindmap", "BSV-SCS3_boq.sql")
mindmap_db_path = os.path.join(base_dir, "kaka", "Mindmap", "mindmap.db")

output_folder = os.path.join(base_dir, "PowerBI_Data")
os.makedirs(output_folder, exist_ok=True)

excel_output_path = os.path.join(base_dir, "KAKA_BIM_PowerBI_Dataset.xlsx")
downloads_excel = r"C:\Users\vophi\Downloads\KAKA_BIM_PowerBI_Dataset.xlsx"

# 2. Extract BOQ elements from BSV-SCS3_boq.sql into in-memory sqlite
mem_conn = sqlite3.connect(":memory:")
with open(boq_sql_path, "r", encoding="utf-8", errors="ignore") as f:
    sql_script = f.read()
mem_conn.executescript(sql_script)
mem_cur = mem_conn.cursor()

# Get BOQ Elements
mem_cur.execute("""
    SELECT element_id, unique_id, category_code, category_name, family_name, 
           type_name, level_name, volume_m3, area_m2, length_m,
           material_name, mark, bsv_phanloai, partition, rebar_weight_kg
    FROM revit_boq_elements
""")
boq_rows = mem_cur.fetchall()
print(f"Loaded {len(boq_rows)} Revit BOQ elements.")

# 3. Extract Schedule 4D from PPMDBSQLite.db
p6_conn = sqlite3.connect(p6_db_path)
p6_cur = p6_conn.cursor()

# Find project
p6_cur.execute("SELECT proj_id, proj_short_name, plan_start_date, plan_end_date FROM project WHERE proj_short_name = 'PRJ-KAKA-BIM'")
prj = p6_cur.fetchone()
if not prj:
    p6_cur.execute("SELECT proj_id, proj_short_name, plan_start_date, plan_end_date FROM project WHERE proj_short_name LIKE '%KAKA%' ORDER BY proj_id DESC LIMIT 1")
    prj = p6_cur.fetchone()

proj_id, proj_code, plan_start, plan_end = prj
print(f"Selected P6 Project: ID={proj_id}, Code={proj_code}, Start={plan_start}, End={plan_end}")

# Get WBS
p6_cur.execute("""
    SELECT wbs_id, wbs_short_name, wbs_name, parent_wbs_id, seq_num
    FROM projwbs
    WHERE proj_id = ?
    ORDER BY seq_num
""", (proj_id,))
wbs_rows = p6_cur.fetchall()

# Map WBS ID to WBS Short Name
wbs_map = {r[0]: (r[1], r[2]) for r in wbs_rows}

# Get Tasks
p6_cur.execute("""
    SELECT task_id, wbs_id, task_code, task_name, target_drtn_hr_cnt, 
           target_start_date, target_end_date, status_code, total_float_hr_cnt
    FROM task
    WHERE proj_id = ?
    ORDER BY target_start_date, task_id
""", (proj_id,))
task_rows = p6_cur.fetchall()
print(f"Loaded {len(task_rows)} P6 tasks.")

# 4. Create Tables for Power BI (Star Schema)

# A. Dim_Project
dim_project = [
    {
        "Project_ID": proj_code,
        "Project_Name": "Dự Án Tổ Hợp KAKA BIM - P6 - CDE",
        "BIM_Manager": "BIM Manager / Director",
        "Plan_Start_Date": plan_start[:10],
        "Plan_End_Date": plan_end[:10],
        "Status": "Đang Triển Khai",
        "Total_BIM_Elements": len(boq_rows),
        "Total_Schedule_Activities": len(task_rows),
        "CDE_Platform": "ACC / Trimble Connect / Dalux",
        "Standard": "ISO 19650 - OpenBIM"
    }
]

# B. Dim_WBS
dim_wbs = []
for r in wbs_rows:
    parent_code = wbs_map.get(r[3], ("ROOT", "Root"))[0] if r[3] else "ROOT"
    discipline = "Kết Cấu (Structural)" if "KC" in r[1] else ("Kiến Trúc (Architectural)" if "KT" in r[1] else "Tổng Thể")
    dim_wbs.append({
        "WBS_ID": r[0],
        "WBS_Code": r[1],
        "WBS_Name": r[2],
        "Discipline": discipline,
        "Parent_WBS_Code": parent_code,
        "Sequence": r[4]
    })

# C. Dim_Components
components_data = [
    ("01.Coc", "Cọc Ly Tâm & Đài Cọc", "Kết Cấu Phần Ngầm", "Foundation & Piles", 10),
    ("02.Mong", "Móng Bê Tông Cốt Thép", "Kết Cấu Phần Ngầm", "Footings / Raft", 12),
    ("03.DaKieng", "Đà Kiềng & Giằng Móng", "Kết Cấu Phần Ngầm", "Tie Beams", 8),
    ("04.Nen", "Nền & Lớp Lót BTXM", "Kết Cấu Tầng Trệt", "Slab on Grade", 15),
    ("05.Cot", "Cột Bê Tông Cốt Thép", "Kết Cấu Thân", "Columns", 10),
    ("06.Vach", "Vách Thang Máy & Vách Cứng", "Kết Cấu Thân", "Shear Walls", 8),
    ("07.Dam", "Dầm Khung Kết Cấu", "Kết Cấu Thân", "Framing Beams", 9),
    ("08.San", "Sàn Sườn & Sàn Nấm BTCT", "Kết Cấu Thân", "Structural Slabs", 12),
    ("09.CauThang", "Cầu Thang Bộ Bộ Môn Kết Cấu", "Giao Thông Đứng", "Stairs", 6),
    ("10.TamCap", "Tam Cấp & Bậc Cửa", "Hạng Mục Phụ", "Entrance Steps", 4),
    ("11.Ramp", "Ramp Dốc Tầng Hầm & Dock", "Giao Thông Ngầm", "Access Ramps", 5),
    ("12.Pit", "Hố Pit Thang Máy & Bể Nước", "Kết Cấu Ngầm Đặc Biệt", "Elevator Pits", 4),
    ("13.Dock", "Sàn Dock Tải Hàng", "Khu Vực Phụ Trợ", "Loading Docks", 5),
    ("14.CauKienPhu", "Cấu Kiện Phụ & Gờ Bê Tông", "Chi Tiết Kiến Trúc/Kết Cấu", "Secondary Works", 4),
]
dim_components = []
for code, name, zone, en_name, duration_std in components_data:
    dim_components.append({
        "Category_Code": code,
        "Category_Name": name,
        "Zone_Group": zone,
        "English_Name": en_name,
        "Standard_Duration_Days": duration_std
    })

# D. Dim_BIM_Elements (1056 elements from Revit)
dim_bim_elements = []
for r in boq_rows:
    elem_id, uid, cat_code, cat_name, fam_name, type_name, lvl, vol, area, length, mat, mark, phanloai, part, rebar_wt = r
    # Cost estimation standard
    # Unit prices: Concrete = 1,450,000 VND/m3; Formwork = 220,000 VND/m2; Rebar = 17,500 VND/kg
    vol_val = vol if vol else 0.0
    area_val = area if area else 0.0
    length_val = length if length else 0.0
    rebar_val = rebar_wt if rebar_wt else (vol_val * 110.0) # default 110kg rebar/m3 concrete
    
    concrete_cost = vol_val * 1450000
    formwork_cost = area_val * 220000
    rebar_cost = rebar_val * 17500
    total_cost = concrete_cost + formwork_cost + rebar_cost

    dim_bim_elements.append({
        "Element_ID": elem_id,
        "Unique_ID": uid,
        "Category_Code": cat_code,
        "Category_Name": cat_name,
        "Family_Name": fam_name,
        "Type_Name": type_name,
        "Level_Name": lvl if lvl else "Ground",
        "Volume_m3": round(vol_val, 3),
        "Area_m2": round(area_val, 2),
        "Length_m": round(length_val, 2),
        "Rebar_Weight_kg": round(rebar_val, 1),
        "Material_Name": mat if mat else "Bê tông cốt thép C30",
        "Mark": mark if mark else "",
        "Estimated_Cost_VND": round(total_cost, 0)
    })

# E. Fact_Schedule_4D
fact_schedule = []
for t in task_rows:
    t_id, w_id, t_code, t_name, drtn_hr, s_date, e_date, status, float_hr = t
    w_info = wbs_map.get(w_id, ("WBS-KC", "Kết Cấu"))
    
    # Calculate days
    drtn_days = round(drtn_hr / 8.0, 1) if drtn_hr else 1.0
    float_days = round(float_hr / 8.0, 1) if float_hr else 0.0
    is_critical = "Có (Đường Găng CPM)" if float_days <= 0 else "Không (Có Float)"
    
    # Extract component code
    cat_code = "00.Chung"
    for comp in components_data:
        if comp[0].split('.')[1] in t_code or comp[0] in t_name or comp[1] in t_name:
            cat_code = comp[0]
            break
            
    # Simulated actual progress %
    # Assign progress based on task dates relative to a baseline
    s_dt = datetime.datetime.strptime(s_date[:10], "%Y-%m-%d")
    cutoff = datetime.datetime(2026, 11, 15)
    if s_dt < cutoff:
        pct_complete = 100 if s_dt < datetime.datetime(2026, 11, 1) else 65
        task_status = "Đã Hoàn Thành" if pct_complete == 100 else "Đang Thi Công"
    else:
        pct_complete = 0
        task_status = "Chưa Bắt Đầu"

    fact_schedule.append({
        "Activity_ID": t_code,
        "Activity_Name": t_name,
        "WBS_Code": w_info[0],
        "Category_Code": cat_code,
        "Start_Date": s_date[:10],
        "End_Date": e_date[:10],
        "Duration_Days": drtn_days,
        "Total_Float_Days": float_days,
        "Is_Critical_Path": is_critical,
        "Status": task_status,
        "Percent_Complete": pct_complete,
        "Primary_Resource": "Đội thi công bê tông & cốt thép KAKA",
        "Cost_Account": "CA-KC-01"
    })

# F. Fact_CDE_Quality (Quality & Coordination tracking)
fact_cde_quality = []
cde_statuses = ["Published (Đã duyệt thi công)", "Shared (Đang phối hợp)", "WIP (Nội bộ thiết kế)"]
contractors = ["Thầu phụ Kết cấu Alpha", "Thầu phụ Cốt thép Delta", "Thầu phụ Ván khuôn Beta", "Đội MEP Gamma"]

for i, comp in enumerate(components_data):
    cat_code, name, zone, _, _ = comp
    total_dwgs = 12 + (i % 5) * 4
    approved_dwgs = int(total_dwgs * (0.85 if i < 6 else 0.65))
    wip_dwgs = total_dwgs - approved_dwgs
    open_issues = (i * 3) % 7
    closed_issues = 15 + i * 2
    
    fact_cde_quality.append({
        "Category_Code": cat_code,
        "Component_Name": name,
        "Subcontractor": contractors[i % len(contractors)],
        "Total_Shop_Drawings": total_dwgs,
        "Drawings_Approved_Published": approved_dwgs,
        "Drawings_In_Review_Shared": int(wip_dwgs * 0.6),
        "Drawings_WIP": wip_dwgs - int(wip_dwgs * 0.6),
        "Approval_Rate_Pct": round(approved_dwgs / total_dwgs * 100, 1),
        "Open_BCF_Issues": open_issues,
        "Resolved_BCF_Issues": closed_issues,
        "Issue_Closure_Rate_Pct": round(closed_issues / (closed_issues + open_issues) * 100, 1),
        "CDE_Platform_Sync": "Đã đồng bộ ACC & Trimble Connect"
    })

print("Generated all in-memory Star Schema tables!")

# 5. Export to CSV files in PowerBI_Data/
tables_to_csv = [
    ("Dim_Project.csv", dim_project),
    ("Dim_WBS.csv", dim_wbs),
    ("Dim_Components.csv", dim_components),
    ("Dim_BIM_Elements.csv", dim_bim_elements),
    ("Fact_Schedule_4D.csv", fact_schedule),
    ("Fact_CDE_Quality.csv", fact_cde_quality),
]

for filename, data in tables_to_csv:
    filepath = os.path.join(output_folder, filename)
    with open(filepath, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=data[0].keys())
        writer.writeheader()
        writer.writerows(data)
    print(f"Exported CSV: {filename} ({len(data)} rows)")

# 6. Export to Excel Master Workbook (KAKA_BIM_PowerBI_Dataset.xlsx)
wb = openpyxl.Workbook()
# remove default sheet
wb.remove(wb.active)

sheet_definitions = [
    ("Dim_Project", dim_project, "0F2B48"),
    ("Dim_WBS", dim_wbs, "1E3A8A"),
    ("Dim_Components", dim_components, "0284C7"),
    ("Fact_Schedule_4D", fact_schedule, "059669"),
    ("Dim_BIM_Elements", dim_bim_elements, "4338CA"),
    ("Fact_CDE_Quality", fact_cde_quality, "D97706"),
]

header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
cell_font = Font(name="Segoe UI", size=9)
thin_border = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)

for title, data, header_color in sheet_definitions:
    ws = wb.create_sheet(title=title)
    ws.views.sheetView[0].showGridLines = True
    
    headers = list(data[0].keys())
    ws.append(headers)
    
    # Style Header
    fill = PatternFill(start_color=header_color, end_color=header_color, fill_type="solid")
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Append Data
    for row_dict in data:
        row_vals = [row_dict[k] for k in headers]
        ws.append(row_vals)
        
    # Format cells
    for row in ws.iter_rows(min_row=2, max_row=len(data) + 1, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.font = cell_font
            cell.border = thin_border
            
    # Auto-fit column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

wb.save(excel_output_path)
print(f"Saved Excel Master: {excel_output_path}")

# Copy to Downloads
import shutil
shutil.copyfile(excel_output_path, downloads_excel)
print(f"Copied Excel Master to: {downloads_excel}")

print("Dataset generation completed successfully!")
